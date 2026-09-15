"""Intake for Eagles Eyrie (RealPage/OneSite-style 'Rent Roll 4.2', flat export).

The source workbook is a 2-property portfolio (sheets 'Eagles Eyrie' and
'Lyndon Crossings'); per the analyst we process the **Eagles Eyrie sheet only**.

Flat one-row-per-unit layout under a header at row 7, closed by an
"Eagles Eyrie Total:" row followed by Status/Charge-Code summaries (skipped).
Every real unit row carries the property look-up code "EAGL" in column A, which
cleanly separates unit rows from the summary rows below.

Columns: A Property Look-Up Code | B Bldg-Unit | C Unit Type | D SQFT
  | E Unit Status | F Resident | G Move-In | H Lease Start | I Lease End
  | J Market Rent (RI) | K Scheduled Charges | L Deposit Held

  * Unit type "NxM ..." -> N BD / M BA; the full string is the floor plan so
    "2x2", "2x2 Townhouse", "3x3 Townhouse" stay distinct.
  * Status: "Occupied No Notice" / "Notice Unrented" -> occupied;
    any "Vacant *" -> vacant (matches the source's Total Occupied 74 / Vacant 6).
  * Market Rent (RI) column present -> real market rents (market_from_inplace off).
  * Per-charge detail is hidden in this export, so "Scheduled Charges" is the
    only per-unit rent figure -> used as the in-place/contract rent.  It bundles
    base Rent + Amenity Rent (property-level: Rent 97,172 + Amenity 3,350 =
    100,522); the split is not available per unit.
"""
import os
import re
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from rent_roll_processor import Unit, RollConfig, build_exhibits  # noqa: E402
from rent_roll_processor.naming import output_filename            # noqa: E402

PROPERTY = "Eagles Eyrie"
SHEET = "Eagles Eyrie"
LOOKUP = "EAGL"                     # property look-up code on every unit row
AS_OF = date(2026, 9, 10)          # source header date 09/10/2026
BDBA_RE = re.compile(r"^(\d+)\s*x\s*(\d+)")


def decode(utype):
    """'2x2 Townhouse' -> ('2x2 Townhouse', 2, 2)."""
    raw = str(utype or "").strip()
    m = BDBA_RE.match(raw)
    if m:
        return raw, int(m.group(1)), int(m.group(2))
    return raw, None, None


def num(v):
    return v if isinstance(v, (int, float)) else None


def load(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)[SHEET]
    units = []
    for r in range(8, ws.max_row + 1):
        if ws.cell(r, 1).value != LOOKUP:              # unit rows only (skip totals/summaries)
            continue
        fp, beds, baths = decode(ws.cell(r, 3).value)
        status = str(ws.cell(r, 5).value or "").strip()
        vacant = status.lower().startswith("vacant")   # Occupied/Notice -> occupied
        name = str(ws.cell(r, 6).value or "").strip()
        sched = num(ws.cell(r, 11).value) or 0
        units.append(Unit.from_dict({
            "unit_id": str(ws.cell(r, 2).value).strip(),
            "unit_type": str(ws.cell(r, 3).value).strip(),   # strict comp key
            "floor_plan": fp,
            "beds": beds,
            "baths": baths,
            "sqft": num(ws.cell(r, 4).value),
            "tenant_name": name or ("Vacant" if vacant else None),
            "occupancy": "Vacant" if vacant else "Occupied",
            "market_rent": num(ws.cell(r, 10).value) or 0,   # J Market Rent (RI)
            "contract_rent": 0 if vacant else sched,         # K Scheduled Charges
            "move_in": ws.cell(r, 7).value,                  # G
            "lease_start": ws.cell(r, 8).value,              # H Lease Start
            "lease_end": ws.cell(r, 9).value,                # I Lease End
        }))
    return units


def main():
    args = sys.argv[1:]
    src = args[0] if args else None
    out = args[1] if len(args) > 1 else output_filename(PROPERTY)
    if not src:
        print("usage: intake_eagles_eyrie.py <source.xlsx> [out.xlsx]")
        return 2
    units = load(src)
    config = RollConfig(property_name=PROPERTY, as_of_date=AS_OF, uw_market_rents={})
    build_exhibits(units, config, out)

    from rent_roll_processor.aggregate import is_occupied
    occ = sum(1 for u in units if is_occupied(u, config.model_occupied))
    vac = sum(1 for u in units if u.occupancy == "Vac")
    print(f"Wrote {out}")
    print(f"  units={len(units)}  occupied={occ}  vacant={vac}  occ%={occ/len(units):.2%}")
    print("  reconcile (mine vs source 'Eagles Eyrie Total:' + Status Summary):")
    for label, mine, src_tot in [
        ("Units", len(units), 80),
        ("Occupied units", occ, 74),
        ("Vacant units", vac, 6),
        ("Sq Ft", sum(u.sqft or 0 for u in units), 112400),
        ("Market Rent", sum(u.market_rent or 0 for u in units), 114490),
        ("Scheduled (in-place)", sum(u.contract_rent or 0 for u in units), 100522),
    ]:
        flag = "OK" if abs(mine - src_tot) < 0.5 else "DIFF"
        print(f"    {label:20s} mine={mine:>12,.2f}  source={src_tot:>12,.2f}  [{flag}]")
    return 0


if __name__ == "__main__":
    sys.exit(main())
