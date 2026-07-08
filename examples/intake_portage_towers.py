"""Intake for Portage Towers (Berkadia 'April' book) — a manual Excel rent roll.

Sheet 'April'.  Header row 3; units from row 4; a 'SOUTH TOWER' banner + repeated
header partway down; grand totals near the bottom (used to reconcile).

Columns: A Name | B Apt | C sq foot | D Unit Type | E Market Rent | F Loss to Lease
  | G Deposit | H Move In | I Lease Expires | J Water & Sewer | K Garage
  | L M-M | M Discount | N Storage | O Gas | P Rent (base) | Q Late fees
  | R Total | S April | T Under/Over | U note | W Occupied/Vacant

Mapping:
  * market_rent   = E (Market Rent)
  * contract_rent = P (base rent);  concession = M (Discount, negative)
      -> net in-place = P + M (netting rule), matching source Loss to Lease.
  * other income  = J Water & Sewer, K Garage, L M-M, N Storage, O Gas, Q Late fees
  * unit type "b-ba W" -> b BD / ba BA - W (wing S/F); "JR" -> unknown BD/BA.
"""
import os
import re
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from rent_roll_processor import Unit, RollConfig, build_exhibits  # noqa: E402
from rent_roll_processor.naming import output_filename            # noqa: E402

SRC = None
PROPERTY = "Portage Towers"
AS_OF = date(2026, 4, 30)

TYPE_RE = re.compile(r"^\s*(\d)\s*-\s*(\d)\s*-?\s*([SF])?\s*$", re.I)
OI_COLS = [("Water & Sewer", 10), ("Garage", 11), ("M-M", 12),
           ("Storage", 14), ("Gas", 15), ("Late fees", 17)]


def decode(dtype):
    raw = str(dtype or "").strip()
    m = TYPE_RE.match(raw)
    if not m:
        # e.g. "JR" / "Jr" -> junior, beds/baths unknown
        return raw.upper(), raw.upper(), None, None
    beds, baths, wing = int(m.group(1)), int(m.group(2)), (m.group(3) or "").upper()
    canon = f"{beds}-{baths}" + (f" {wing}" if wing else "")
    fp = f"{beds} BD / {baths} BA" + (f" - {wing}" if wing else "")
    return canon, fp, beds, baths


def num(v):
    return v if isinstance(v, (int, float)) else None


def load(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)["April"]
    units = []
    tower = "North"
    for r in range(4, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and "SOUTH TOWER" in a.upper():
            tower = "South"
            continue
        sqft = ws.cell(r, 3).value
        dtype = ws.cell(r, 4).value
        # unit row: has a square footage number AND a unit type
        if not isinstance(sqft, (int, float)) or dtype in (None, ""):
            continue
        code, fp, beds, baths = decode(dtype)
        status = str(ws.cell(r, 23).value or "").strip()
        vacant = status.lower() == "vacant"
        name = str(a).strip() if a else ""
        base = num(ws.cell(r, 16).value) or 0        # P
        disc = num(ws.cell(r, 13).value) or 0        # M
        mm = num(ws.cell(r, 12).value) or 0          # L (month-to-month fee)
        oi = {label: num(ws.cell(r, c).value) for label, c in OI_COLS
              if num(ws.cell(r, c).value)}
        units.append(Unit.from_dict({
            "unit_id": str(ws.cell(r, 2).value).strip(),
            "unit_type": code,
            "floor_plan": fp,
            "beds": beds,
            "baths": baths,
            "sqft": sqft,
            "tenant_name": name or ("Vacant" if vacant else None),
            "occupancy": "Vacant" if vacant else "Occupied",
            "market_rent": num(ws.cell(r, 5).value) or 0,     # E
            "contract_rent": 0 if vacant else base,           # P (base)
            "concession": 0 if vacant else disc,              # M (discount)
            "other_income": sum(oi.values()),
            "other_income_items": oi,
            "mtm": "MTM" if mm else None,
            "move_in": ws.cell(r, 8).value,
            "lease_start": ws.cell(r, 8).value,   # no lease-sign date -> Move-In fallback
            "lease_end": ws.cell(r, 9).value,
            "designation": (str(ws.cell(r, 21).value).strip()
                            if ws.cell(r, 21).value else None) or tower,
        }))
    return units


def main():
    global SRC
    args = sys.argv[1:]
    SRC = args[0] if args else None
    out = args[1] if len(args) > 1 else output_filename(PROPERTY)
    if not SRC:
        print("usage: intake_portage_towers.py <source.xlsx> [out.xlsx]")
        return 2
    units = load(SRC)
    config = RollConfig(property_name=PROPERTY, as_of_date=AS_OF, uw_market_rents={})
    build_exhibits(units, config, out)

    from rent_roll_processor.aggregate import is_occupied
    occ = sum(1 for u in units if is_occupied(u, config.model_occupied))
    vac = sum(1 for u in units if u.occupancy == "Vac")
    print(f"Wrote {out}")
    print(f"  units={len(units)}  occupied={occ}  vacant={vac}  occ%={occ/len(units):.2%}")
    # reconcile against the source grand-totals row
    print("  reconcile (mine vs source row 388):")
    checks = [
        ("Sq Ft", sum(u.sqft or 0 for u in units), 332650),
        ("Market Rent (E)", sum(u.market_rent for u in units), 421990),
        ("Base Rent (P)", sum(u.contract_rent for u in units), 375130),
        ("Discount (M)", sum(u.concession for u in units), -20680),
        ("Other income", sum(u.other_income for u in units), 15080 + 4830 + 1875 + 1020 + 12705 - 268.05),
    ]
    for label, mine, src in checks:
        flag = "OK" if abs(mine - src) < 1 else "DIFF"
        print(f"    {label:16s} mine={mine:>12,.2f}  source={src:>12,.2f}  [{flag}]")
    return 0


if __name__ == "__main__":
    sys.exit(main())
