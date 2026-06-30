"""Intake / normalization for the Station J Town (Yardi) rent roll.

Source layout (sheet 'Report1'):
    row 6  : headers  Unit | Unit Type | Unit Sqft | Name | Market Rent |
                      Actual Rent | Resident Deposit | Other Deposits |
                      Move In | Lease Expiration | Move Out
    row 7  : section banner 'Current/Notice/Vacant Residents'
    rows 8-391 : unit rows
    rows 392+  : subtotals / summary groups

Floor-plan codes decode as  stjt<beds><letter><baths>[R]:
    2A1 -> 2 BD / 1 BA - A,  2B1 -> 2 BD / 1 BA - B,  2B2 -> 2 BD / 2 BA - B,
    2C1 -> 2 BD / 1 BA - C,  2C2 -> 2 BD / 2 BA - C.   Trailing 'R' = renovated.

Notes / assumptions (flagged to the analyst):
  * Source has no lease-sign date.  Per the SOP fallback, lease_start is
    proxied as Lease Expiration - 12 months (only counts toward Recent Leases
    when that lands on/before the as-of date).
  * Source has no other-income column, so Other Income = 0.
  * Underwriting market rents intentionally left blank (LTL columns blank).
"""
import os
import re
import sys
from datetime import date

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from rent_roll_processor import Unit, RollConfig, build_exhibits  # noqa: E402

SRC = sys.argv[1] if len(sys.argv) > 1 else None
OUT = sys.argv[2] if len(sys.argv) > 2 else "Rent_Roll_Exhibits_Station_J_Town_05.22.26.xlsx"

CODE_RE = re.compile(r"^stjt(\d)([A-Z])(\d)$", re.IGNORECASE)


def minus_12_months(d):
    if d is None:
        return None
    try:
        return d.replace(year=d.year - 1)
    except ValueError:  # Feb 29 -> Feb 28
        return d.replace(year=d.year - 1, day=28)


def decode(code):
    """stjt2B1R -> (floor_plan, beds, baths, renovated)."""
    code = code.strip()
    reno = code.upper().endswith("R")
    base = code[:-1] if reno else code
    m = CODE_RE.match(base)
    if not m:
        return code, None, None, reno
    beds, letter, baths = int(m.group(1)), m.group(2), int(m.group(3))
    return f"{beds} BD / {baths} BA - {letter}", beds, baths, reno


def clean_name(raw):
    if raw is None:
        return None
    n = str(raw).strip().rstrip("#*$").strip()
    return n or None


def is_unit_row(a):
    return a and isinstance(a, str) and re.match(r"^[A-Z]?\d", a.strip())


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Report1"]
    units = []
    for r in range(8, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if not is_unit_row(a):
            continue
        code = str(ws.cell(r, 2).value or "").strip()
        fp, beds, baths, reno = decode(code)
        name_raw = clean_name(ws.cell(r, 4).value)
        nl = (name_raw or "").lower()
        if nl == "vacant":
            occ, tenant = "Vacant", "VACANT"
        elif nl == "model":
            occ, tenant = "Model", "MODEL"
        else:
            occ, tenant = "Occupied", name_raw
        lease_exp = ws.cell(r, 10).value
        units.append(Unit.from_dict({
            "unit_id": a.strip(),
            "unit_type": code.replace("stjt", ""),
            "floor_plan": fp,
            "sqft": ws.cell(r, 3).value,
            "beds": beds,
            "baths": baths,
            "renovated": "Yes" if reno else None,
            "reno_type": "Full" if reno else None,
            "tenant_name": tenant,
            "occupancy": occ,
            "market_rent": ws.cell(r, 5).value,
            "contract_rent": ws.cell(r, 6).value,
            "other_income": 0,
            "move_in": ws.cell(r, 9).value,
            "lease_start": minus_12_months(lease_exp),   # proxy (see module docstring)
            "lease_end": lease_exp,
        }))
    return units


def main():
    if not SRC:
        print("usage: intake_station_jtown.py <source.xlsx> [out.xlsx]")
        return 2
    units = load(SRC)
    config = RollConfig(
        property_name="Station J Town",
        as_of_date=date(2026, 5, 22),
        uw_market_rents={},   # left blank per request -> LTL columns blank
    )
    build_exhibits(units, config, OUT)

    occ = sum(1 for u in units if u.is_occupied)
    vac = sum(1 for u in units if u.is_vacant)
    print(f"Wrote {OUT}")
    print(f"  units={len(units)}  occupied={occ}  vacant={vac}  occ%={occ/len(units):.2%}")
    from collections import Counter
    fp = Counter(u.floor_plan for u in units)
    for k in sorted(fp):
        print(f"    {k}: {fp[k]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
