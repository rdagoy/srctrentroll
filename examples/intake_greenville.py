"""Intake for the Greenville Portfolio (AppFolio 'Rent Roll', 4 properties).

Sheet 'Sheet1'.  A portfolio export: a "Property Groups: Greenville Portfolio"
header, then one **property banner** per building, its unit rows, a per-property
subtotal row ("<n> Units  <pct> Occupied"), and a grand-total row at the bottom
("Total  78 Units ...").  Unit numbers repeat across buildings, so per the
portfolio rule (PROCESS.md Step 0) each unit is scoped to its property:

  * unit #    -> "<CODE>-<raw unit>"  (e.g. MLK-4, VAN-12)  -- unique portfolio-wide
  * unit type -> the per-property CODE (the source carries no unit type, so the
                 property IS the grouping key); floor plan = the building label.

Columns: A Tenant | B Unit | C Status | D Rent | E Deposit | F Lease From
  | G Lease To | H Move-in | I Move-out

The source has no unit-type / SF / bed-bath / market-rent columns, so
market rent = in-place Rent (placeholder) and SF/beds/baths are left blank
(unit mix groups by property).  Commercial units (ids starting "Comm") are
flagged retail and dropped by exclude_retail (#29).
"""
import os
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from rent_roll_processor import Unit, RollConfig, build_exhibits  # noqa: E402
from rent_roll_processor.naming import output_filename            # noqa: E402

PROPERTY = "Greenville Portfolio"
AS_OF = date(2026, 6, 30)          # source header "As of: 06/30/2026"

# (banner substring, short code, building label) -- scopes unit #/type per property.
PROPS = [
    ("80-82 MLK",    "MLK", "80-82 MLK Dr"),
    ("Van Nostrand", "VAN", "190 Van Nostrand"),
    ("Dwight",       "DWT", "200 Dwight St"),
    ("Stegman",      "STG", "150 Stegman St"),
]


def num(v):
    return v if isinstance(v, (int, float)) else None


def match_prop(banner):
    for sub, code, label in PROPS:
        if sub.lower() in banner.lower():
            return code, label
    return None, None


def load(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)["Sheet1"]
    units = []
    code = label = None
    for r in range(11, ws.max_row + 1):
        a = ws.cell(r, 1).value          # Tenant (blank for vacant)
        b = ws.cell(r, 2).value          # Unit (or "<n> Units" on subtotal rows)
        c = ws.cell(r, 3).value          # Status (or "<pct> Occupied" on subtotals)

        # Property banner: tenant text, no unit / status.
        if isinstance(a, str) and a.strip() and (b in (None, "")) and (c in (None, "")):
            code, label = match_prop(a)
            continue
        # Skip subtotal / grand-total / blank rows.
        if b in (None, ""):
            continue
        bs = str(b).strip()
        if "unit" in bs.lower() or (isinstance(a, str) and a.strip().lower() == "total"):
            continue
        if isinstance(c, str) and "occupied" in c.lower():
            continue
        if code is None:
            continue

        status = str(c or "").strip()
        vacant = status.lower().startswith("vacant")   # Vacant-Unrented / Vacant-Rented
        name = str(a or "").strip()
        rent = num(ws.cell(r, 4).value) or 0
        retail = bs.lower().startswith("comm")          # Comm1 / Comm2 -> retail (#29)

        units.append(Unit.from_dict({
            "unit_id": f"{code}-{bs}",                  # scoped unique per property
            "unit_type": code,                          # property code = grouping key
            "floor_plan": label,                        # building label
            "is_retail": retail,
            "beds": None, "baths": None, "sqft": None,  # not in source
            "tenant_name": name or ("Vacant" if vacant else None),
            "occupancy": "Vacant" if vacant else "Occupied",
            "market_rent": 0 if vacant else rent,       # no market col -> in-place placeholder
            "contract_rent": 0 if vacant else rent,
            "lease_start": ws.cell(r, 6).value or ws.cell(r, 8).value,  # Lease From, else Move-in
            "lease_end": ws.cell(r, 7).value,
            "move_in": ws.cell(r, 8).value,
            "move_out": ws.cell(r, 9).value,
        }))
    return units


def main():
    args = sys.argv[1:]
    src = args[0] if args else None
    out = args[1] if len(args) > 1 else output_filename(PROPERTY)
    if not src:
        print("usage: intake_greenville.py <source.xlsx> [out.xlsx]")
        return 2
    units = load(src)
    config = RollConfig(property_name=PROPERTY, as_of_date=AS_OF, uw_market_rents={})
    build_exhibits(units, config, out)

    from rent_roll_processor.aggregate import is_occupied
    from rent_roll_processor.schema import is_retail_unit
    print(f"Wrote {out}")
    print(f"  parsed units={len(units)}  (incl. commercial)")

    # per-property breakdown (residential, i.e. what lands in the exhibits)
    res = [u for u in units if not is_retail_unit(u)]
    for _, code, lbl in PROPS:
        pu = [u for u in res if u.unit_type == code]
        occ = sum(1 for u in pu if is_occupied(u, config.model_occupied))
        print(f"    {code} {lbl:18s} units={len(pu):2d} occ={occ:2d} "
              f"rent={sum(u.contract_rent for u in pu):>10,.2f}")
    r_occ = sum(1 for u in res if is_occupied(u, config.model_occupied))
    print(f"  residential (in exhibits): units={len(res)} occ={r_occ} "
          f"occ%={r_occ/len(res):.2%}  (excluded {len(units)-len(res)} retail)")

    print("  reconcile (parsed vs source grand-total row):")
    for label, mine, src_tot in [
        ("Units", len(units), 78),
        ("Rent", sum(u.contract_rent for u in units), 93870.80),
    ]:
        flag = "OK" if abs(mine - src_tot) < 0.5 else "DIFF"
        print(f"    {label:10s} mine={mine:>12,.2f}  source={src_tot:>12,.2f}  [{flag}]")
    return 0


if __name__ == "__main__":
    sys.exit(main())
