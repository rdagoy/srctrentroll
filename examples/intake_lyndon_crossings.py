"""Intake for Lyndon Crossings (RealPage/OneSite 'Rent Roll', flat export).

Standalone single-property file (same PM format as the Eagles Eyrie sheet, but a
slightly different column layout -- this export adds an "Expected Move-Out"
column, shifting Market Rent to K and Scheduled Charges to L).

Flat one-row-per-unit layout under a header at row 7, closed by a
"Lyndon Crossings Total:" row, then Status/Charge-Code summaries and a separate
"Future Resident Details" block (different columns) -- all skipped.  Every real
unit row carries "Lyndon Crossings" in column A (the Future block does not),
which cleanly selects the unit rows.

Columns: A Property | B Bldg-Unit | C Unit Type | D SQFT | E Unit Status
  | F Resident | G Move-In | H Lease Start | I Lease End | J Expected Move-Out
  | K Market Rent | L Scheduled Charges | M Deposit Held

  * Unit type "NxM-<plan>" -> N BD / M BA; full string is the floor plan so
    "2x2-Parkwood" and "2x2-Gentry" stay distinct.
  * Status: "Occupied No Notice" / "Notice Unrented" -> occupied;
    any "Vacant *" -> vacant (matches source Total Occupied 55 / Vacant 5).
  * Market Rent column present -> real market rents (market_from_inplace off).
  * Per-charge detail hidden -> "Scheduled Charges" is the only per-unit rent
    figure -> used as in-place.  It bundles base Rent 61,373 + Month-to-Month
    Rent 1,299 + Amenity Rent 1,850 = 64,522 (property level; no per-unit split).
"""
import os
import re
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from rent_roll_processor import Unit, RollConfig, build_exhibits  # noqa: E402
from rent_roll_processor.naming import output_filename            # noqa: E402

PROPERTY = "Lyndon Crossings"
SHEET = "Lyndon Crossings"
AS_OF = date(2026, 8, 31)          # source period "Aug 2026" / file 2026.08.31
BDBA_RE = re.compile(r"^(\d+)\s*x\s*(\d+)")


def decode(utype):
    """'2x2-Parkwood' -> ('2x2-Parkwood', 2, 2)."""
    raw = str(utype or "").strip()
    m = BDBA_RE.match(raw)
    if m:
        return raw, int(m.group(1)), int(m.group(2))
    return raw, None, None


def num(v):
    return v if isinstance(v, (int, float)) else None


def load(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)[SHEET]
    units = []
    for r in range(8, ws.max_row + 1):
        if ws.cell(r, 1).value != PROPERTY:            # unit rows only (skip totals/summaries/future)
            continue
        fp, beds, baths = decode(ws.cell(r, 3).value)
        status = str(ws.cell(r, 5).value or "").strip()
        vacant = status.lower().startswith("vacant")   # Occupied/Notice -> occupied
        name = str(ws.cell(r, 6).value or "").strip()
        sched = num(ws.cell(r, 12).value) or 0
        units.append(Unit.from_dict({
            "unit_id": str(ws.cell(r, 2).value).strip(),
            "unit_type": str(ws.cell(r, 3).value).strip(),   # strict comp key
            "floor_plan": fp,
            "beds": beds,
            "baths": baths,
            "sqft": num(ws.cell(r, 4).value),
            "tenant_name": name or ("Vacant" if vacant else None),
            "occupancy": "Vacant" if vacant else "Occupied",
            "market_rent": num(ws.cell(r, 11).value) or 0,   # K Market Rent
            "contract_rent": 0 if vacant else sched,         # L Scheduled Charges
            "move_in": ws.cell(r, 7).value,                  # G
            "lease_start": ws.cell(r, 8).value,              # H Lease Start
            "lease_end": ws.cell(r, 9).value,                # I Lease End
            "move_out": ws.cell(r, 10).value,                # J Expected Move-Out
        }))
    return units


def main():
    args = sys.argv[1:]
    src = args[0] if args else None
    out = args[1] if len(args) > 1 else output_filename(PROPERTY)
    if not src:
        print("usage: intake_lyndon_crossings.py <source.xlsx> [out.xlsx]")
        return 2
    units = load(src)
    config = RollConfig(property_name=PROPERTY, as_of_date=AS_OF, uw_market_rents={})
    build_exhibits(units, config, out, source_path=src, source_sheet=SHEET)

    from rent_roll_processor.aggregate import is_occupied
    occ = sum(1 for u in units if is_occupied(u, config.model_occupied))
    vac = sum(1 for u in units if u.occupancy == "Vac")
    print(f"Wrote {out}")
    print(f"  units={len(units)}  occupied={occ}  vacant={vac}  occ%={occ/len(units):.2%}")
    print("  reconcile (mine vs source 'Lyndon Crossings Total:' + Status Summary):")
    for label, mine, src_tot in [
        ("Units", len(units), 60),
        ("Occupied units", occ, 55),
        ("Vacant units", vac, 5),
        ("Sq Ft", sum(u.sqft or 0 for u in units), 55760),
        ("Market Rent", sum(u.market_rent or 0 for u in units), 71703),
        ("Scheduled (in-place)", sum(u.contract_rent or 0 for u in units), 64522),
    ]:
        flag = "OK" if abs(mine - src_tot) < 0.5 else "DIFF"
        print(f"    {label:20s} mine={mine:>12,.2f}  source={src_tot:>12,.2f}  [{flag}]")
    return 0


if __name__ == "__main__":
    sys.exit(main())
