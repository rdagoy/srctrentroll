"""Intake for Pendale Towers (AppFolio 'Rent Roll', single property).

Sheet 'Sheet1'.  A single property ("Pendale Towers - 460 Washington Road,
Pittsburgh PA") -- not a portfolio.  Flat unit list under a 16-column header
(row 10), closed by a grand-total row ("Total 132 Units").  Same AppFolio
layout as Nova, but this export DOES carry a Market Rent column.

Columns: A Unit | B Tags | C BD/BA | D Tenant | E Status | F Sqft
  | G Market Rent | H Rent | I Deposit | J Lease From | K Lease To
  | L Move-in | M Move-out | ...

Classification decisions (flagged in the run summary):
  * Status: Current / Notice-Unrented -> occupied; Vacant-Unrented AND
    Vacant-Rented -> vacant.  Gives 126/132 = 95.5%, matching the source's
    own "95.5% Occupied" grand-total figure (a vacant-rented unit is
    physically vacant now, with a future lease).
  * Two non-residential rows carry BD/BA "--/--": the "Rooftop Antenna"
    (Pinnacle Towers LLC cell-site license) and the "106/107 Commercial
    Suite".  Both flagged is_retail -> excluded from the residential exhibits.
  * "Guest Suite (3rd Floor)" / tenant "Birgo AirBnB" is an owner-operated
    short-term-rental unit carrying $0 in-place rent.  Kept as an occupied
    residential unit; the default $0-rent placeholder rule marks its in-place
    to its $1,025 market rent (treat-as-market-if-converted).  Its $0 status is
    surfaced in the run summary so it can instead be carved out as non-revenue
    if the STR income is underwritten separately.
"""
import os
import re
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from rent_roll_processor import Unit, RollConfig, build_exhibits  # noqa: E402
from rent_roll_processor.naming import output_filename            # noqa: E402

PROPERTY = "Pendale Towers"
AS_OF = date(2026, 7, 30)          # source header "As of: 07/30/2026"
BDBA_RE = re.compile(r"^(\d+)\s*/\s*([\d.]+)")


def decode(bb):
    """'1/1.00' -> ('1 BD / 1 BA', 1, 1.0, residential=True); '--/--' -> non-residential."""
    raw = str(bb or "").strip()
    m = BDBA_RE.match(raw)
    if not m:
        return raw, None, None, False
    beds = int(m.group(1))
    baths = float(m.group(2))
    ba_disp = str(int(baths)) if baths == int(baths) else str(baths)
    return f"{beds} BD / {ba_disp} BA", beds, baths, True


def num(v):
    return v if isinstance(v, (int, float)) else None


def load(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)["Sheet1"]
    units = []
    for r in range(12, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if not (isinstance(a, str) and a.strip()):
            continue
        if a.startswith("Pendale"):
            continue                                    # per-property banner row
        status = str(ws.cell(r, 5).value or "").strip()
        if "occupied" in status.lower():
            continue                                    # grand-total row(s)
        bb = ws.cell(r, 3).value
        if bb is None and not status:
            continue                                    # stray/blank row
        fp, beds, baths, residential = decode(bb)
        name = str(ws.cell(r, 4).value or "").strip()
        vacant = status.lower().startswith("vacant")    # incl. Vacant-Rented -> vacant
        mkt = num(ws.cell(r, 7).value) or 0
        rent = num(ws.cell(r, 8).value) or 0

        occ = "Vacant" if vacant else "Occupied"

        units.append(Unit.from_dict({
            "unit_id": a.strip(),
            "unit_type": fp,                            # BD/BA grouping / comp key
            "floor_plan": fp,
            "beds": beds,
            "baths": baths,
            "sqft": num(ws.cell(r, 6).value),
            "tenant_name": name or ("Vacant" if vacant else None),
            "occupancy": occ,
            "market_rent": mkt,
            "contract_rent": 0 if vacant else rent,
            "lease_start": ws.cell(r, 10).value or ws.cell(r, 12).value,   # Lease From, else Move-in
            "lease_end": ws.cell(r, 11).value,          # Lease To
            "is_retail": not residential,               # '--/--' -> non-residential
        }))
    return units


def main():
    args = sys.argv[1:]
    src = args[0] if args else None
    out = args[1] if len(args) > 1 else output_filename(PROPERTY)
    if not src:
        print("usage: intake_pendale.py <source.xlsx> [out.xlsx]")
        return 2
    units = load(src)
    config = RollConfig(property_name=PROPERTY, as_of_date=AS_OF, uw_market_rents={})
    build_exhibits(units, config, out)

    from rent_roll_processor.aggregate import is_occupied
    from rent_roll_processor.schema import is_retail_unit
    res = [u for u in units if not is_retail_unit(u)]
    occ = sum(1 for u in res if is_occupied(u, config.model_occupied))
    vac = sum(1 for u in res if u.occupancy == "Vac")
    zero = [u for u in res if is_occupied(u, config.model_occupied) and not (u.contract_rent or 0)]
    print(f"Wrote {out}")
    print(f"  parsed units={len(units)} (incl. {len(units)-len(res)} non-residential excluded)")
    print(f"  residential={len(res)}  occupied={occ}  vacant={vac}  occ%={occ/len(res):.2%}")
    if zero:
        print(f"  occupied $0-rent units marked to market (placeholder): "
              f"{', '.join(u.unit_id for u in zero)}")
    print("  reconcile (ALL parsed rows vs source grand-total row):")
    for label, mine, src_tot in [
        ("Units", len(units), 132),
        ("Sq Ft", sum(u.sqft or 0 for u in units), 114814),
        ("Market Rent", sum(u.market_rent or 0 for u in units), 165945),
        ("Rent (in-place)", sum(u.contract_rent or 0 for u in units), 159144.16),
    ]:
        flag = "OK" if abs(mine - src_tot) < 0.5 else "DIFF"
        print(f"    {label:16s} mine={mine:>13,.2f}  source={src_tot:>13,.2f}  [{flag}]")
    return 0


if __name__ == "__main__":
    sys.exit(main())
