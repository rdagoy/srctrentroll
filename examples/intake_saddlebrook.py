"""Intake for Saddlebrook at Tates Creek (Yardi 'Rent Roll with Lease Charges').

Sheet 'Report1'.  Same charge-ledger layout as 20 Grand Ave: one main row per
unit (market rent + first charge code), then one sub-row per additional charge
code, closed by a per-unit "Total" row.  Sections are banner rows:

  * "Current/Notice/Vacant Residents"  <- the as-of snapshot (parsed)
  * "Future Residents/Applicants"      <- upcoming leases for units already
                                          counted; SKIPPED (would double-count)
  * "Summary Groups" / charge summary  <- SKIPPED (used only to reconcile)

Columns: A Unit | B Unit Type | C Unit Sq Ft | D Resident | E Name
  | F Market Rent | G Charge Code | H Amount | I Resident Deposit
  | J Other Deposit | K Move In | L Lease Expiration | M Move Out | N Balance

Charge-code mapping (validated against the source's current-only charge summary,
total 236,899).  Per #32, the exact source charge code is used verbatim as the
OneLineRR column header for every other-income / concession / employee-discount
line item (no renaming):
  * BASE (contract rent) = conrent
  * EMPLOYEE DISCOUNT     = conempl  (negative; kept separate, NOT netted)
  * OTHER INCOME          = conpest, contfinc, coninsu, congarg, conpetrt
      (pet -> leftmost), conliab
  * no concession codes present.

Unit types are plan codes ("stc_a1".."stc_b4", "stc_eh") with no embedded
BD/BA.  Beds are inferred from the prefix -- "a*" = 1BR, "b*" = 2BR (the SF gap,
731-987 vs 1204-1405, supports it); baths are left blank for the analyst to
confirm in the OneLineRR Checking table.  "stc_eh" (unit EH, 3,400 SF) is a
large single home kept as its own type.  Unit 092 (tenant "MODEL") is the one
non-revenue unit (auto-detected -> netted to $0, out of occ/vac).
No lease-sign date in the source -> Move In is the lease-start placeholder.
"""
import os
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from rent_roll_processor import Unit, RollConfig, build_exhibits  # noqa: E402
from rent_roll_processor.naming import output_filename            # noqa: E402

PROPERTY = "Saddlebrook at Tates Creek"
AS_OF = date(2026, 7, 5)           # source header "As Of = 07/05/2026"

BASE = {"conrent"}
EMP = {"conempl"}                  # employee discount (negative)
BANNERS = {"Current/Notice/Vacant Residents", "Future Residents/Applicants",
           "Summary Groups"}


def decode(utype):
    """Plan code -> (floor_plan, beds, baths).  Beds from a/b prefix; baths TBD."""
    u = str(utype or "").strip().lower()
    if u.startswith("stc_a"):
        return "1 Bedroom", 1, None
    if u.startswith("stc_b"):
        return "2 Bedroom", 2, None
    if u == "stc_eh":
        return "Executive Home", None, None
    return str(utype or "").strip(), None, None


def num(v):
    return v if isinstance(v, (int, float)) else None


def load(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)["Report1"]

    def is_banner(r):
        a = ws.cell(r, 1).value
        return isinstance(a, str) and a.strip() in BANNERS

    def is_unit_start(r):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        return (isinstance(a, str) and a.strip() and not is_banner(r)
                and isinstance(b, str) and b.strip())

    stops = sorted({r for r in range(5, ws.max_row + 1)
                    if is_banner(r) or is_unit_start(r)} | {ws.max_row + 1})

    units, section = [], None
    for r in range(5, ws.max_row + 1):
        if is_banner(r):
            section = ws.cell(r, 1).value.strip()
            continue
        if section != "Current/Notice/Vacant Residents" or not is_unit_start(r):
            continue
        end = min(x for x in stops if x > r)          # next unit-start or banner
        charges = {}
        for rr in range(r, end):
            g, h = ws.cell(rr, 7).value, ws.cell(rr, 8).value
            if isinstance(g, str) and g not in ("Charge", "Total") and isinstance(h, (int, float)):
                charges[g] = charges.get(g, 0) + h

        fp, beds, baths = decode(ws.cell(r, 2).value)
        name = str(ws.cell(r, 5).value or "").strip()
        resid = str(ws.cell(r, 4).value or "").strip()
        vacant = name.upper() == "VACANT" or resid.upper() == "VACANT" or not name

        contract = sum(v for k, v in charges.items() if k in BASE)
        # exact source charge codes as line-item labels (#32)
        emp_items = {k: v for k, v in charges.items() if k in EMP}
        oi_items = {k: v for k, v in charges.items() if k not in BASE and k not in EMP}
        emp = sum(emp_items.values())

        units.append(Unit.from_dict({
            "unit_id": str(ws.cell(r, 1).value).strip(),
            "unit_type": str(ws.cell(r, 2).value).strip(),   # strict comp key
            "floor_plan": fp,
            "beds": beds,
            "baths": baths,
            "sqft": num(ws.cell(r, 3).value),
            "tenant_name": name or ("Vacant" if vacant else None),
            "occupancy": "Vacant" if vacant else "Occupied",  # "MODEL" -> auto non-rev
            "market_rent": num(ws.cell(r, 6).value) or 0,     # F
            "contract_rent": 0 if vacant else contract,       # sum(BASE)
            "emp_discount": 0 if vacant else emp,             # conempl (negative)
            "emp_discount_items": {} if vacant else emp_items,
            "other_income": sum(oi_items.values()),
            "other_income_items": oi_items,
            "move_in": ws.cell(r, 11).value,
            "lease_start": ws.cell(r, 11).value,   # no sign date -> Move In placeholder
            "lease_end": ws.cell(r, 12).value,     # Lease Expiration
            "move_out": ws.cell(r, 13).value,
        }))
    return units


def main():
    args = sys.argv[1:]
    src = args[0] if args else None
    out = args[1] if len(args) > 1 else output_filename(PROPERTY)
    if not src:
        print("usage: intake_saddlebrook.py <source.xlsx> [out.xlsx]")
        return 2
    units = load(src)
    config = RollConfig(property_name=PROPERTY, as_of_date=AS_OF, uw_market_rents={})
    build_exhibits(units, config, out)

    from rent_roll_processor.aggregate import is_occupied
    is_model = lambda u: (u.tenant_name or "").strip().lower() == "model"
    model = sum(1 for u in units if is_model(u))
    occ = sum(1 for u in units if is_occupied(u, config.model_occupied) and not is_model(u))
    vac = sum(1 for u in units if u.occupancy == "Vac")
    print(f"Wrote {out}")
    print(f"  units(parsed)={len(units)}  occupied={occ}  vacant={vac}  "
          f"non-rev(model)={model}  occ%={occ/len(units):.2%} (occ/total, model in denom)")

    def sf_of(pred):
        return sum(u.sqft or 0 for u in units if pred(u))
    print("  reconcile (parsed vs source Summary Groups / charge summary):")
    contract = sum(u.contract_rent for u in units)
    other = sum(u.other_income for u in units)
    emp = sum(u.emp_discount for u in units)
    for label, mine, src_tot in [
        ("Units", len(units), 181),
        ("Vacant units", vac, 9),
        ("Non-rev units", model, 1),
        ("Total Sq Ft", sf_of(lambda u: True), 185350),
        ("Occupied Sq Ft", sf_of(lambda u: is_occupied(u, False) and not is_model(u)), 174790),
        ("Vacant Sq Ft", sf_of(lambda u: u.occupancy == "Vac"), 9573),
        ("Non-rev Sq Ft", sf_of(is_model), 987),
        ("Market Rent (all)", sum(u.market_rent for u in units), 244146),
        ("Base rent (conrent)", contract, 229667),
        ("Other income", other, 10324),
        ("Employee disc", emp, -3092),
        ("Lease charges", contract + other + emp, 236899),
    ]:
        flag = "OK" if abs(mine - src_tot) < 0.5 else "DIFF"
        print(f"    {label:20s} mine={mine:>13,.2f}  source={src_tot:>13,.2f}  [{flag}]")
    return 0


if __name__ == "__main__":
    sys.exit(main())
