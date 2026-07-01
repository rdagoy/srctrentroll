"""Fidelity test: regenerate the Heritage Hill exhibits from the extracted
unit data and diff every value against the original reference workbook.

This is the regression guard for the engine.  Run:

    python -m tests.validate_heritage  [path/to/reference.xlsm]

Exit code is non-zero if any value mismatches.
"""
import json
import os
import sys
from datetime import date, datetime

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from rent_roll_processor import Unit, RollConfig, build_exhibits  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
UNITS_JSON = os.path.join(REPO, "examples", "heritage_hill_units.json")

# Source -> Unit field mapping for the extracted reference JSON.
FIELD_MAP = {
    "plan": "floor_plan", "sqft": "sqft", "bed": "beds", "bath": "baths",
    "mkt": "market_rent", "contract": "contract_rent", "unit": "unit_id",
    "utype": "unit_type", "prop": "property_name", "desig": "designation",
    "renod": "renovated", "renotype": "reno_type", "name": "tenant_name",
    "ltype": "lease_type", "renostat": "reno_status", "occ": "occupancy",
    "mtm": "mtm", "renew": "renew_status", "asking": "asking_rent",
    "conc": "concession", "empdisc": "emp_discount", "other": "other_income",
    "movein": "move_in", "lstart": "lease_start", "lend": "lease_end",
}

UW_MARKET_RENTS = {
    "1 BD / 1 BA": 1250,
    "2 BD / 1 BA": 1375,
    "2 BD / 1.5 BA": 1375,
    "2 BD / 2 BA - A": 1500,
    "2 BD / 2 BA - B": 1500,
}


def load_units():
    raw = json.load(open(UNITS_JSON))
    units = []
    for row in raw:
        d = {FIELD_MAP[k]: v for k, v in row.items() if k in FIELD_MAP}
        units.append(Unit.from_dict(d))
    return units


def approx(a, b, tol=1e-6):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= max(tol, abs(b) * 1e-9)
    if isinstance(a, datetime):
        a = a.date()
    if isinstance(b, datetime):
        b = b.date()
    return a == b


def norm(v):
    if isinstance(v, datetime):
        return v.date()
    if v == "":
        return None
    return v


def compare_sheet(gen_ws, ref_ws, rows, cols, label, mismatches):
    for r in rows:
        for c in cols:
            g = norm(gen_ws.cell(r, c).value)
            ref = norm(ref_ws.cell(r, c).value)
            if g is None and ref is None:
                continue
            if not approx(g, ref):
                mismatches.append(f"{label} {openpyxl.utils.get_column_letter(c)}{r}: gen={g!r} ref={ref!r}")


def main():
    ref_path = sys.argv[1] if len(sys.argv) > 1 else os.environ.get("REF_XLSM")
    units = load_units()
    config = RollConfig(
        property_name="Heritage Hill Estates",
        as_of_date=date(2026, 1, 31),
        uw_market_rents=UW_MARKET_RENTS,
        model_occupied=True,        # the reference exhibits fold the model into occupied
        derive_vacant_market=False,  # reference is a frozen fixture (its own vacant rents)
        expand_nonrev=False,         # its model already carries the rent/concession offset
    )
    out = os.path.join(HERE, "_heritage_generated.xlsx")
    build_exhibits(units, config, out)
    print(f"Generated {out}")

    if not ref_path or not os.path.exists(ref_path):
        print("No reference file given; generation succeeded. "
              "Pass the reference .xlsm path to run the value diff.")
        return 0

    gen = openpyxl.load_workbook(out, data_only=False)
    ref = openpyxl.load_workbook(ref_path, data_only=False)
    mismatches = []

    # Unit Mix & RR Summary: data rows 6-11 + stat block 13-22, cols C..S
    compare_sheet(gen["Unit Mix & RR Summary"], ref["Unit Mix & RR Summary"],
                  list(range(5, 23)), range(3, 20), "UMRS", mismatches)
    # Recent Leases: rows 6-12, cols C..W
    compare_sheet(gen["Recent Leases"], ref["Recent Leases"],
                  list(range(6, 13)), range(3, 24), "RL", mismatches)
    # Bed mix: rows 5-19, cols C..S
    compare_sheet(gen["Unit Mix based on # of Beds"], ref["Unit Mix based on # of Beds"],
                  list(range(5, 20)), range(3, 20), "BED", mismatches)
    # Pres. Rent Roll values: rows 8-128, cols D..AG (skip styling-only cols)
    compare_sheet(gen["Pres. Rent Roll"], ref["Pres. Rent Roll"],
                  list(range(8, 129)), range(4, 34), "PR", mismatches)

    if mismatches:
        print(f"\n{len(mismatches)} MISMATCHES:")
        for m in mismatches[:80]:
            print("  ", m)
        return 1
    print("\nALL VALUES MATCH ✓")
    return 0


if __name__ == "__main__":
    sys.exit(main())
